"""Template 7.1 rendering: branded cover, body pages, hours table and grilling annexes (PDF + DOCX)."""

import io
import uuid
from datetime import datetime

import pypdfium2 as pdfium
from docx import Document as DocxDocument
from PIL import Image

from scopeprofit.domain.branding import DEFAULT_BRANDING, Branding
from scopeprofit.render.template71 import DocumentContext, corner_png, render_docx, render_markdown, render_pdf
from tests.fixtures.yard import FRANPRZDEV_BRANDING, yard_brief

PRICING_WORDS = ("Presupuesto", "Costo", "USD", "INVERSIÓN", "Precio por hora")


def context(branding: Branding = FRANPRZDEV_BRANDING) -> DocumentContext:
    return DocumentContext(
        "Sistema de Gestión Integral para Desarmadero",
        "Desarmadero Ejemplo",
        "pro@example.com",
        datetime(2026, 10, 14),
        2,
        yard_brief(uuid.uuid4()),
        branding,
    )


def page_texts(pdf: bytes) -> list[str]:
    document = pdfium.PdfDocument(pdf)
    return [document[index].get_textpage().get_text_range() for index in range(len(document))]


def hex_rgb(value: str) -> tuple[int, int, int]:
    raw = value.removeprefix("#")
    return int(raw[0:2], 16), int(raw[2:4], 16), int(raw[4:6], 16)


def close(pixel: tuple[int, ...], expected: tuple[int, int, int], tolerance: int = 12) -> bool:
    return all(abs(channel - target) <= tolerance for channel, target in zip(pixel[:3], expected, strict=True))


def test_pdf_cover_and_body_follow_the_branded_layout() -> None:
    pdf = render_pdf(context())
    texts = page_texts(pdf)
    assert len(texts) >= 3
    cover = " ".join(texts[0].split())
    for expected in (
        "FranPrzDev",
        "Octubre 2026",
        "ACUERDO DE",
        "DESARROLLO",
        "Análisis, diseño y desarrollo de software",
    ):
        assert expected in cover
    assert "DDev" not in cover
    body = " ".join(" ".join(texts[1:]).split())
    for expected in (
        "Documento de alcance: Sistema de Gestión Integral para Desarmadero",
        "1. Resumen ejecutivo",
        "6. Estimación orientativa en horas",
        "Etapa",
        "Actividades principales",
        "TOTAL - Estimación orientativa del alcance",
        "180–240 hs",
        "Anexo A. Glosario",
        "Anexo B. Decisiones",
        "Anexo C. Preguntas resueltas",
    ):
        assert expected in body
    for word in PRICING_WORDS:
        assert word not in body


def test_pdf_cover_renders_the_brand_colors() -> None:
    pdf = render_pdf(context())
    page = pdfium.PdfDocument(pdf)[0]
    image: Image.Image = page.render(scale=1).to_pil().convert("RGB")
    width, height = image.size
    primary = hex_rgb(FRANPRZDEV_BRANDING.primary_color)
    accent = hex_rgb(FRANPRZDEV_BRANDING.accent_color)
    assert close(image.getpixel((width - 5, height - 5)), primary)
    assert close(image.getpixel((width - 5, height - 70)), accent)
    background = image.getpixel((width // 2, int(height * 0.3)))
    assert min(background[:3]) > 200, "body background stays light behind the pattern"


def test_default_branding_is_neutral_scopeprofit() -> None:
    cover = " ".join(page_texts(render_pdf(context(DEFAULT_BRANDING)))[0].split())
    assert "ScopeProfit" in cover
    assert "DOCUMENTO DE" in cover
    assert "FranPrzDev" not in cover


def test_docx_has_cover_header_hours_table_and_annexes() -> None:
    document = DocxDocument(io.BytesIO(render_docx(context())))
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    assert "ACUERDO DE" in text
    assert "DESARROLLO" in text
    assert "Análisis, diseño y desarrollo de software" in text
    assert "Anexo A. Glosario" in text
    header = document.sections[1].header.tables[0].rows[0].cells
    assert header[0].text == "FranPrzDev"
    assert header[1].text == "Octubre 2026"
    hours = next(table for table in document.tables if table.rows[0].cells[0].text == "Etapa")
    assert [cell.text for cell in hours.rows[0].cells] == ["Etapa", "Actividades principales", "Horas"]
    assert hours.rows[-1].cells[0].text == "TOTAL - Estimación orientativa del alcance"
    assert hours.rows[-1].cells[-1].text == "180–240 hs"
    assert not any(word in text for word in PRICING_WORDS)


def test_markdown_lists_the_grilling_annexes() -> None:
    markdown = render_markdown(context())
    assert "## Anexo A. Glosario\n- **Vehículo**:" in markdown
    assert "- D-01 Bajas lógicas en todo el sistema (aceptada):" in markdown
    assert "- Q-01 (ronda 1):" in markdown


def test_corner_png_uses_the_brand_colors() -> None:
    image = Image.open(io.BytesIO(corner_png(FRANPRZDEV_BRANDING))).convert("RGBA")
    assert image.getpixel((image.width - 2, image.height - 2))[:3] == hex_rgb(FRANPRZDEV_BRANDING.primary_color)
    assert image.getpixel((image.width - 2, 2))[:3] == hex_rgb(FRANPRZDEV_BRANDING.accent_color)
