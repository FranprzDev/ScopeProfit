"""Template 7.1: the fixed scope document rendered as Markdown, PDF and DOCX.

The document language is Spanish (the product's document language since v1). The PDF and DOCX follow a
branded proposal layout: a full cover page (brand and month top, a two-weight title, a brand block with
the tagline, diagonal corner bands and a faint background pattern) followed by clean body pages with the
seven numbered sections, an hours table (Etapa | Actividades principales | Horas, with a TOTAL row) and
annexes for the glossary, decision records and resolved questions captured while grilling the scope.
Brand name, tagline, colors and cover title come from the owner's `Branding`. There is no pricing.
"""

import io
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from xml.sax.saxutils import escape

from docx import Document as DocxDocument
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor
from PIL import Image, ImageDraw
from reportlab.lib import colors
from reportlab.lib.enums import TA_JUSTIFY
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.pdfmetrics import stringWidth
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen.canvas import Canvas
from reportlab.platypus import (
    BaseDocTemplate,
    Frame,
    KeepTogether,
    ListFlowable,
    ListItem,
    NextPageTemplate,
    PageBreak,
    PageTemplate,
    Paragraph,
    Spacer,
    Table,
    TableStyle,
)
from reportlab.platypus.flowables import Flowable

from scopeprofit.domain.brief import BriefData
from scopeprofit.domain.branding import DEFAULT_BRANDING, Branding
from scopeprofit.enums import ExportFormat

PENDING = "Pendiente de definición"
DISCLAIMER = "Rango orientativo sujeto a los supuestos e incertidumbres indicados; no es un precio."
MONTHS = (
    "Enero",
    "Febrero",
    "Marzo",
    "Abril",
    "Mayo",
    "Junio",
    "Julio",
    "Agosto",
    "Septiembre",
    "Octubre",
    "Noviembre",
    "Diciembre",
)
SCOPE_LABELS = ("Incluidos", "Excluidos", "Supuestos", "Criterios de aceptación")
TEXT_COLOR = "#2B2B2B"
MUTED_COLOR = "#6B7280"
GRID_COLOR = "#D5D9DE"
HEADER_FILL = "#EEF1F4"
PAGE_FILL = "#FBFBFC"


@dataclass(frozen=True)
class DocumentContext:
    project_name: str
    client_name: str
    author: str
    date: datetime
    version: int
    brief: BriefData
    branding: Branding = field(default=DEFAULT_BRANDING)


def _hours(value: float) -> str:
    return str(int(value)) if value == int(value) else f"{value:g}"


def _range(minimum: float, maximum: float) -> str:
    return _hours(minimum) if minimum == maximum else f"{_hours(minimum)}–{_hours(maximum)}"


def _bullets(values: list[str]) -> str:
    return "\n".join(f"- {value}" for value in values) or f"- {PENDING}"


def _yes_no(value: bool) -> str:
    return "Sí" if value else "No"


def month_year(date: datetime) -> str:
    return f"{MONTHS[date.month - 1]} {date.year}"


def _scope_lists(brief: BriefData) -> tuple[tuple[str, list[str]], ...]:
    values = (brief.included, brief.excluded, brief.assumptions, brief.acceptance_criteria)
    return tuple(zip(SCOPE_LABELS, values, strict=True))


def _estimate_rows(brief: BriefData) -> list[list[str]]:
    return [
        [e.module, e.activities or e.uncertainty, f"{_range(e.min_hours, e.max_hours)} hs"] for e in brief.estimates
    ]


def _total_label(brief: BriefData) -> str:
    total_min, total_max = brief.total_hours()
    return f"{_range(total_min, total_max)} hs"


# ----- Markdown ----------------------------------------------------------------------------------


def _markdown_annexes(brief: BriefData) -> str:
    parts: list[str] = []
    if brief.glossary:
        terms = "\n".join(
            f"- **{t.term}**: {t.definition}" + (f" _Evitar_: {', '.join(t.avoid)}" if t.avoid else "")
            for t in brief.glossary
        )
        parts.append(f"## Anexo A. Glosario\n{terms}\n")
    if brief.decisions:
        decisions = "\n".join(f"- {d.id} {d.title} ({d.status.value}): {d.decision}" for d in brief.decisions)
        parts.append(f"## Anexo B. Decisiones\n{decisions}\n")
    if brief.resolved_questions:
        resolved = "\n".join(
            f"- {q.id} (ronda {q.round}): {q.question} → {q.answer}" for q in brief.resolved_questions
        )
        parts.append(f"## Anexo C. Preguntas resueltas\n{resolved}\n")
    return "\n".join(parts)


def render_markdown(doc: DocumentContext) -> str:
    brief = doc.brief
    requirements = "\n".join(
        f'- {r.id} | {r.description} | {r.type.value} | {r.priority.value} | Fuente: "{r.source}" | {r.system_note}'
        for r in brief.requirements
    )
    questions = "\n".join(
        f"- {q.id}: {q.question} | {q.reason} | Bloquea estimación: {_yes_no(q.blocks_estimate)}"
        for q in brief.questions
    )
    risks = "\n".join(
        f"- {r.id}: {r.description} | Impacto: {r.impact} | Mitigación/supuesto: {r.mitigation} | "
        f"Semáforo: {r.severity.value}"
        for r in brief.risks
    )
    estimates = "\n".join(
        f"- {e.module}: {_hours(e.min_hours)}–{_hours(e.max_hours)} h | {e.uncertainty}" for e in brief.estimates
    )
    total_min, total_max = brief.total_hours()
    annexes = _markdown_annexes(brief)
    return (
        f"# {doc.project_name}\n"
        f"Cliente: {doc.client_name}\n"
        f"Fecha: {doc.date.isoformat()}\n"
        f"Versión: 0.{doc.version}\n"
        f"Autor: {doc.author}\n\n"
        f"## 1. Resumen ejecutivo\n{brief.summary}\n\n"
        f"## 2. Requerimientos identificados\n{requirements}\n\n"
        f"## 3. Preguntas pendientes\n{questions or 'Sin preguntas pendientes'}\n\n"
        f"## 4. Riesgos y ambigüedades\n{risks or 'Sin riesgos identificados'}\n\n"
        f"## 5. Alcance\n### Incluidos\n{_bullets(brief.included)}\n### Excluidos\n{_bullets(brief.excluded)}\n"
        f"### Supuestos\n{_bullets(brief.assumptions)}\n"
        f"### Criterios de aceptación\n{_bullets(brief.acceptance_criteria)}\n\n"
        f"## 6. Estimación orientativa en horas\n{estimates}\n"
        f"Total: {_hours(total_min)}–{_hours(total_max)} h. {DISCLAIMER}\n\n"
        f"## 7. Próximos pasos\n{_bullets(brief.next_steps)}\n" + (f"\n{annexes}" if annexes else "")
    )


# ----- PDF ---------------------------------------------------------------------------------------

FONT_DIR = Path(__file__).resolve().parent / "fonts"
FONT_LIGHT = "Poppins-Light"
FONT_REGULAR = "Poppins-Regular"
FONT_MEDIUM = "Poppins-Medium"
FONT_SEMIBOLD = "Poppins-SemiBold"
FONT_BOLD = "Poppins-Bold"
_FONTS = (FONT_LIGHT, FONT_REGULAR, FONT_MEDIUM, FONT_SEMIBOLD, FONT_BOLD)
_fonts_registered = False


def register_fonts() -> None:
    """Poppins (SIL OFL, bundled in `fonts/`) covers Spanish accents and typographic dashes."""
    global _fonts_registered
    if _fonts_registered:
        return
    for name in _FONTS:
        pdfmetrics.registerFont(TTFont(name, str(FONT_DIR / f"{name}.ttf")))
    pdfmetrics.registerFontFamily(FONT_REGULAR, normal=FONT_REGULAR, bold=FONT_SEMIBOLD)
    _fonts_registered = True


PAGE_WIDTH, PAGE_HEIGHT = A4
MARGIN_X = 20 * mm
BODY_TOP = 34 * mm
BODY_BOTTOM = 34 * mm


def _p(text: str, style: ParagraphStyle) -> Paragraph:
    return Paragraph(escape(text).replace("\n", "<br/>"), style)


def _draw_pattern(canvas: Canvas, branding: Branding) -> None:
    """Off-white page with a faint lattice of large arcs, like a woven watermark."""
    canvas.setFillColor(colors.HexColor(PAGE_FILL))
    canvas.rect(0, 0, PAGE_WIDTH, PAGE_HEIGHT, stroke=0, fill=1)
    canvas.saveState()
    canvas.setStrokeColor(colors.HexColor(branding.primary_color))
    canvas.setStrokeAlpha(0.022)
    canvas.setLineWidth(11)
    radius = 150 * mm
    step = 60 * mm
    for column in range(-2, 6):
        for row in range(-1, 6):
            canvas.circle(column * step, PAGE_HEIGHT - row * step * 1.3, radius, stroke=1, fill=0)
    canvas.restoreState()


def _draw_corner(canvas: Canvas, branding: Branding) -> None:
    """Accent and primary diagonal bands in the bottom-right corner."""
    canvas.saveState()
    accent = canvas.beginPath()
    accent.moveTo(PAGE_WIDTH * 0.72, 34)
    accent.lineTo(PAGE_WIDTH, 34)
    accent.lineTo(PAGE_WIDTH, 112)
    accent.lineTo(PAGE_WIDTH * 0.86, 112)
    accent.close()
    canvas.setFillColor(colors.HexColor(branding.accent_color))
    canvas.drawPath(accent, stroke=0, fill=1)
    primary = canvas.beginPath()
    primary.moveTo(PAGE_WIDTH * 0.63, 0)
    primary.lineTo(PAGE_WIDTH, 0)
    primary.lineTo(PAGE_WIDTH, 26)
    primary.lineTo(PAGE_WIDTH * 0.665, 26)
    primary.close()
    canvas.setFillColor(colors.HexColor(branding.primary_color))
    canvas.drawPath(primary, stroke=0, fill=1)
    canvas.restoreState()


def _fit(text: str, font: str, size: float, width: float) -> float:
    while size > 14 and stringWidth(text, font, size) > width:
        size -= 1
    return size


@dataclass(frozen=True)
class _Pages:
    doc: DocumentContext

    def cover(self, canvas: Canvas, _template: BaseDocTemplate) -> None:
        branding = self.doc.branding
        primary = colors.HexColor(branding.primary_color)
        accent = colors.HexColor(branding.accent_color)
        _draw_pattern(canvas, branding)
        _draw_corner(canvas, branding)
        top = PAGE_HEIGHT - 26 * mm
        canvas.setFillColor(accent)
        canvas.setFont(FONT_SEMIBOLD, 15)
        canvas.drawString(MARGIN_X, top, branding.brand_name)
        canvas.setFillColor(primary)
        canvas.setFont(FONT_SEMIBOLD, 15)
        canvas.drawRightString(PAGE_WIDTH - MARGIN_X, top, month_year(self.doc.date))
        width = PAGE_WIDTH - 2 * MARGIN_X
        light_size = _fit(branding.cover_title_light, FONT_LIGHT, 54, width)
        bold_size = _fit(branding.cover_title_bold, FONT_BOLD, 58, width)
        middle = PAGE_HEIGHT * 0.56
        canvas.setFont(FONT_LIGHT, light_size)
        canvas.drawString(MARGIN_X - 2, middle, branding.cover_title_light)
        canvas.setFont(FONT_BOLD, bold_size)
        canvas.drawString(MARGIN_X - 2, middle - bold_size * 1.12, branding.cover_title_bold)
        canvas.setFillColor(colors.HexColor(MUTED_COLOR))
        canvas.setFont(FONT_REGULAR, 12)
        subtitle = f"{self.doc.project_name} · {self.doc.client_name} · Versión 0.{self.doc.version}"
        canvas.drawString(MARGIN_X, middle - bold_size * 1.12 - 30, subtitle[:110])
        block = 44 * mm
        canvas.setFillColor(primary)
        canvas.setFont(FONT_SEMIBOLD, 19)
        canvas.drawString(MARGIN_X, block, branding.brand_name)
        canvas.setFillColor(accent)
        canvas.setFont(FONT_LIGHT, 13.5)
        canvas.drawString(MARGIN_X, block - 21, branding.tagline)
        canvas.setStrokeColor(accent)
        canvas.setLineWidth(1.2)
        canvas.line(MARGIN_X, block - 38, MARGIN_X + 30 * mm, block - 38)

    def body(self, canvas: Canvas, _template: BaseDocTemplate) -> None:
        branding = self.doc.branding
        _draw_pattern(canvas, branding)
        _draw_corner(canvas, branding)
        top = PAGE_HEIGHT - 24 * mm
        canvas.setFillColor(colors.HexColor(branding.primary_color))
        canvas.setFont(FONT_SEMIBOLD, 15)
        canvas.drawString(MARGIN_X, top, branding.brand_name)
        canvas.drawRightString(PAGE_WIDTH - MARGIN_X, top, month_year(self.doc.date))
        canvas.setFillColor(colors.HexColor(MUTED_COLOR))
        canvas.setFont(FONT_REGULAR, 8)
        canvas.drawString(MARGIN_X, 14 * mm, f"Documento de alcance — plantilla 7.1 · página {canvas.getPageNumber()}")


@dataclass(frozen=True)
class _Styles:
    body: ParagraphStyle
    cell: ParagraphStyle
    cell_head: ParagraphStyle
    total: ParagraphStyle
    title: ParagraphStyle
    meta: ParagraphStyle
    h2: ParagraphStyle
    h3: ParagraphStyle


def _styles(branding: Branding) -> _Styles:
    text = colors.HexColor(TEXT_COLOR)
    body = ParagraphStyle(
        "body", fontName=FONT_REGULAR, fontSize=9.6, leading=14.2, textColor=text, alignment=TA_JUSTIFY, spaceAfter=5
    )
    cell = ParagraphStyle("cell", parent=body, fontSize=8.6, leading=12, alignment=0, spaceAfter=0)
    return _Styles(
        body=body,
        cell=cell,
        cell_head=ParagraphStyle("cell_head", parent=cell, fontName=FONT_MEDIUM),
        total=ParagraphStyle("total", parent=cell, textColor=colors.HexColor(branding.accent_color)),
        title=ParagraphStyle(
            "title",
            fontName=FONT_SEMIBOLD,
            fontSize=17,
            leading=24,
            textColor=colors.HexColor(branding.accent_color),
            spaceAfter=4,
        ),
        meta=ParagraphStyle("meta", parent=body, fontSize=8.6, leading=12, textColor=colors.HexColor(MUTED_COLOR)),
        h2=ParagraphStyle(
            "h2",
            fontName=FONT_SEMIBOLD,
            fontSize=13.5,
            leading=18,
            textColor=colors.HexColor(branding.primary_color),
            spaceBefore=14,
            spaceAfter=5,
            keepWithNext=1,
        ),
        h3=ParagraphStyle(
            "h3",
            fontName=FONT_SEMIBOLD,
            fontSize=10.2,
            leading=14,
            textColor=colors.HexColor(branding.accent_color),
            spaceBefore=6,
            spaceAfter=3,
            keepWithNext=1,
        ),
    )


def _grid(data: list[list[Flowable]], widths: list[float], *, total_row: bool = False) -> Table:
    table = Table(data, colWidths=widths, repeatRows=1)
    commands: list[tuple[object, ...]] = [
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(HEADER_FILL)),
        ("BACKGROUND", (0, 1), (-1, -1), colors.white),
        ("GRID", (0, 0), (-1, -1), 0.6, colors.HexColor(GRID_COLOR)),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("TOPPADDING", (0, 0), (-1, -1), 7),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 7),
        ("LEFTPADDING", (0, 0), (-1, -1), 7),
        ("RIGHTPADDING", (0, 0), (-1, -1), 7),
    ]
    if total_row:
        commands.append(("SPAN", (0, -1), (-2, -1)))
    table.setStyle(TableStyle(commands))
    return table


def _table(header: list[str], rows: list[list[str]], styles: _Styles, widths: list[float]) -> Table:
    data = [[_p(value, styles.cell_head) for value in header], *[[_p(v, styles.cell) for v in row] for row in rows]]
    return _grid(data, widths)


def _hours_table(brief: BriefData, styles: _Styles, width: float) -> Table:
    rows: list[list[Flowable]] = [
        [_p(value, styles.cell_head) for value in ("Etapa", "Actividades principales", "Horas")],
        *[[_p(value, styles.cell) for value in row] for row in _estimate_rows(brief)],
        [
            _p("TOTAL - Estimación orientativa del alcance", styles.total),
            _p("", styles.total),
            _p(_total_label(brief), styles.total),
        ],
    ]
    return _grid(rows, [w * width for w in (0.26, 0.58, 0.16)], total_row=True)


def _pdf_story(doc: DocumentContext, styles: _Styles, width: float) -> list[Flowable]:
    brief = doc.brief

    def bullets(values: list[str]) -> ListFlowable:
        items = [ListItem(_p(value, styles.body), leftIndent=14) for value in values] or [
            ListItem(_p(PENDING, styles.body), leftIndent=14)
        ]
        return ListFlowable(items, bulletType="bullet", start="•", leftIndent=14, bulletFontName=FONT_REGULAR)

    story: list[Flowable] = [
        NextPageTemplate("body"),
        PageBreak(),
        Paragraph(f"<u>Documento de alcance: {escape(doc.project_name)}</u>", styles.title),
        _p(
            f"Cliente: {doc.client_name} · Versión 0.{doc.version} · Fecha: {doc.date.date().isoformat()} · "
            f"Autor: {doc.author}",
            styles.meta,
        ),
        _p("1. Resumen ejecutivo", styles.h2),
        _p(brief.summary or PENDING, styles.body),
        _p("2. Requerimientos identificados", styles.h2),
    ]
    if brief.requirements:
        story.append(
            _table(
                ["ID", "Descripción", "Tipo", "Prioridad", "Fuente", "Nota"],
                [
                    [r.id, r.description, r.type.value, r.priority.value, f"«{r.source}»", r.system_note]
                    for r in brief.requirements
                ],
                styles,
                [w * width for w in (0.08, 0.27, 0.13, 0.13, 0.22, 0.17)],
            )
        )
    else:
        story.append(_p(PENDING, styles.body))
    story.append(_p("3. Preguntas pendientes", styles.h2))
    if brief.questions:
        story.append(
            _table(
                ["ID", "Pregunta", "Motivo", "¿Bloquea estimación?"],
                [[q.id, q.question, q.reason, _yes_no(q.blocks_estimate)] for q in brief.questions],
                styles,
                [w * width for w in (0.09, 0.43, 0.32, 0.16)],
            )
        )
    else:
        story.append(_p("Sin preguntas pendientes", styles.body))
    story.append(_p("4. Riesgos y ambigüedades", styles.h2))
    if brief.risks:
        story.append(
            _table(
                ["ID", "Riesgo", "Impacto", "Mitigación / supuesto", "Semáforo"],
                [[r.id, r.description, r.impact, r.mitigation, r.severity.value] for r in brief.risks],
                styles,
                [w * width for w in (0.09, 0.3, 0.22, 0.27, 0.12)],
            )
        )
    else:
        story.append(_p("Sin riesgos identificados", styles.body))
    story.append(_p("5. Alcance", styles.h2))
    for label, values in _scope_lists(brief):
        story.extend([_p(label, styles.h3), bullets(values)])
    story.append(_p("6. Estimación orientativa en horas", styles.h2))
    if brief.estimates:
        story.extend([_hours_table(brief, styles, width), Spacer(1, 6)])
    else:
        story.append(_p(PENDING, styles.body))
    total_min, total_max = brief.total_hours()
    story.extend(
        [
            _p(f"Total: {_hours(total_min)}–{_hours(total_max)} h. {DISCLAIMER}", styles.meta),
            _p("7. Próximos pasos", styles.h2),
            bullets(brief.next_steps),
        ]
    )
    if brief.glossary:
        story.append(_p("Anexo A. Glosario", styles.h2))
        for term in brief.glossary:
            avoid = f" <i>Evitar: {escape(', '.join(term.avoid))}</i>" if term.avoid else ""
            story.append(
                Paragraph(f"<b>{escape(term.term)}</b>: {escape(term.definition)}{avoid}", styles.body),
            )
    if brief.decisions:
        story.append(_p("Anexo B. Decisiones", styles.h2))
        for decision in brief.decisions:
            parts: list[Flowable] = [
                _p(f"{decision.id}. {decision.title} ({decision.status.value})", styles.h3),
                _p(decision.decision, styles.body),
            ]
            if decision.considered_options:
                parts.append(_p("Opciones consideradas: " + "; ".join(decision.considered_options), styles.meta))
            if decision.consequences:
                parts.append(_p(f"Consecuencias: {decision.consequences}", styles.meta))
            story.append(KeepTogether(parts))
    if brief.resolved_questions:
        story.append(_p("Anexo C. Preguntas resueltas", styles.h2))
        story.append(
            _table(
                ["Ronda", "Pregunta", "Respuesta"],
                [[str(q.round), f"{q.id}. {q.question}", q.answer] for q in brief.resolved_questions],
                styles,
                [w * width for w in (0.1, 0.45, 0.45)],
            )
        )
    return story


def render_pdf(doc: DocumentContext) -> bytes:
    register_fonts()
    pages = _Pages(doc)
    width = PAGE_WIDTH - 2 * MARGIN_X
    buffer = io.BytesIO()
    template = BaseDocTemplate(
        buffer,
        pagesize=A4,
        title=f"{doc.branding.cover_title_light} {doc.branding.cover_title_bold} — {doc.project_name}",
        author=doc.author,
        subject="Documento de alcance — plantilla 7.1",
        creator=doc.branding.brand_name,
    )
    cover_frame = Frame(MARGIN_X, BODY_BOTTOM, width, PAGE_HEIGHT - BODY_TOP - BODY_BOTTOM, id="cover")
    body_frame = Frame(MARGIN_X, BODY_BOTTOM, width, PAGE_HEIGHT - BODY_TOP - BODY_BOTTOM - 8 * mm, id="body")
    template.addPageTemplates(
        [
            PageTemplate(id="cover", frames=[cover_frame], onPage=pages.cover),
            PageTemplate(id="body", frames=[body_frame], onPage=pages.body),
        ]
    )
    template.build(_pdf_story(doc, _styles(doc.branding), width))
    return buffer.getvalue()


# ----- DOCX --------------------------------------------------------------------------------------


def _rgb(hex_color: str) -> RGBColor:
    return RGBColor.from_string(hex_color.removeprefix("#").upper())


def corner_png(branding: Branding, width: int = 900, height: int = 300) -> bytes:
    """The PDF's corner bands as a transparent PNG, for the DOCX cover."""
    image = Image.new("RGBA", (width, height), (0, 0, 0, 0))
    draw = ImageDraw.Draw(image)
    scale = height / 112
    draw.polygon(
        [(width * 0.25, height - 34 * scale), (width, height - 34 * scale), (width, 0), (width * 0.62, 0)],
        fill=branding.accent_color,
    )
    draw.polygon(
        [(0, height), (width, height), (width, height - 26 * scale), (width * 0.1, height - 26 * scale)],
        fill=branding.primary_color,
    )
    buffer = io.BytesIO()
    image.save(buffer, format="PNG")
    return buffer.getvalue()


def _shade(cell_element: OxmlElement, fill: str) -> None:
    properties = cell_element.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), fill.removeprefix("#"))
    properties.append(shading)


def render_docx(doc: DocumentContext) -> bytes:
    branding = doc.branding
    document = DocxDocument()
    document.core_properties.title = doc.project_name
    document.core_properties.author = doc.author
    document.core_properties.subject = "Documento de alcance — plantilla 7.1"
    normal = document.styles["Normal"]
    normal.font.name = "Poppins"
    normal.font.size = Pt(10)
    normal.font.color.rgb = _rgb(TEXT_COLOR)
    for name, size, color in (
        ("Title", 18, branding.accent_color),
        ("Heading 1", 14, branding.primary_color),
        ("Heading 2", 11, branding.accent_color),
    ):
        style = document.styles[name]
        style.font.name = "Poppins"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = _rgb(color)
    section = document.sections[0]
    section.page_width, section.page_height = Mm(210), Mm(297)
    section.left_margin = section.right_margin = Mm(20)
    section.top_margin = section.bottom_margin = Mm(22)

    def brand_row(container_table_parent: DocxDocument, left_color: str) -> None:
        row = container_table_parent.add_table(rows=1, cols=2)
        row.alignment = WD_TABLE_ALIGNMENT.CENTER
        left = row.rows[0].cells[0].paragraphs[0].add_run(branding.brand_name)
        left.bold = True
        left.font.size = Pt(14)
        left.font.color.rgb = _rgb(left_color)
        right_paragraph = row.rows[0].cells[1].paragraphs[0]
        right_paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        right = right_paragraph.add_run(month_year(doc.date))
        right.bold = True
        right.font.size = Pt(14)
        right.font.color.rgb = _rgb(branding.primary_color)

    brand_row(document, branding.accent_color)
    spacer = document.add_paragraph()
    spacer.paragraph_format.space_before = Pt(190)
    light = spacer.add_run(branding.cover_title_light)
    light.font.name = "Poppins Light"
    light.font.size = Pt(44)
    light.font.color.rgb = _rgb(branding.primary_color)
    light.add_break(WD_BREAK.LINE)
    bold = spacer.add_run(branding.cover_title_bold)
    bold.bold = True
    bold.font.size = Pt(48)
    bold.font.color.rgb = _rgb(branding.primary_color)
    subtitle = document.add_paragraph().add_run(
        f"{doc.project_name} · {doc.client_name} · Versión 0.{doc.version}"
    )
    subtitle.font.color.rgb = _rgb(MUTED_COLOR)
    block = document.add_paragraph()
    block.paragraph_format.space_before = Pt(170)
    name = block.add_run(branding.brand_name)
    name.bold = True
    name.font.size = Pt(16)
    name.font.color.rgb = _rgb(branding.primary_color)
    name.add_break(WD_BREAK.LINE)
    tagline = block.add_run(branding.tagline)
    tagline.font.size = Pt(12)
    tagline.font.color.rgb = _rgb(branding.accent_color)
    corner = document.add_paragraph()
    corner.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    corner.add_run().add_picture(io.BytesIO(corner_png(branding)), width=Mm(70))

    body = document.add_section(WD_SECTION.NEW_PAGE)
    body.header.is_linked_to_previous = False
    header_table = body.header.add_table(rows=1, cols=2, width=Mm(170))
    header_left = header_table.rows[0].cells[0].paragraphs[0].add_run(branding.brand_name)
    header_left.bold = True
    header_left.font.color.rgb = _rgb(branding.primary_color)
    header_right_paragraph = header_table.rows[0].cells[1].paragraphs[0]
    header_right_paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header_right = header_right_paragraph.add_run(month_year(doc.date))
    header_right.bold = True
    header_right.font.color.rgb = _rgb(branding.primary_color)

    brief = doc.brief
    document.add_heading(f"Documento de alcance: {doc.project_name}", level=0)
    document.add_paragraph(
        f"Cliente: {doc.client_name} · Versión 0.{doc.version} · Fecha: {doc.date.date().isoformat()} · "
        f"Autor: {doc.author}"
    )

    def table(header: list[str], rows: list[list[str]], total: str | None = None) -> None:
        grid = document.add_table(rows=1, cols=len(header))
        grid.style = "Table Grid"
        for index, label in enumerate(header):
            cell = grid.rows[0].cells[index]
            cell.paragraphs[0].add_run(label).bold = True
            _shade(cell._tc, HEADER_FILL)
        for row in rows:
            cells = grid.add_row().cells
            for index, value in enumerate(row):
                cells[index].text = value
        if total is not None:
            cells = grid.add_row().cells
            merged = cells[0].merge(cells[-2])
            label_run = merged.paragraphs[0].add_run("TOTAL - Estimación orientativa del alcance")
            label_run.font.color.rgb = _rgb(branding.accent_color)
            total_run = cells[-1].paragraphs[0].add_run(total)
            total_run.font.color.rgb = _rgb(branding.accent_color)

    def bullets(values: list[str]) -> None:
        for value in values or [PENDING]:
            document.add_paragraph(value, style="List Bullet")

    document.add_heading("1. Resumen ejecutivo", level=1)
    document.add_paragraph(brief.summary or PENDING)
    document.add_heading("2. Requerimientos identificados", level=1)
    if brief.requirements:
        table(
            ["ID", "Descripción", "Tipo", "Prioridad", "Fuente", "Nota"],
            [
                [r.id, r.description, r.type.value, r.priority.value, f"«{r.source}»", r.system_note]
                for r in brief.requirements
            ],
        )
    else:
        document.add_paragraph(PENDING)
    document.add_heading("3. Preguntas pendientes", level=1)
    if brief.questions:
        table(
            ["ID", "Pregunta", "Motivo", "¿Bloquea estimación?"],
            [[q.id, q.question, q.reason, _yes_no(q.blocks_estimate)] for q in brief.questions],
        )
    else:
        document.add_paragraph("Sin preguntas pendientes")
    document.add_heading("4. Riesgos y ambigüedades", level=1)
    if brief.risks:
        table(
            ["ID", "Riesgo", "Impacto", "Mitigación / supuesto", "Semáforo"],
            [[r.id, r.description, r.impact, r.mitigation, r.severity.value] for r in brief.risks],
        )
    else:
        document.add_paragraph("Sin riesgos identificados")
    document.add_heading("5. Alcance", level=1)
    for label, values in _scope_lists(brief):
        document.add_heading(label, level=2)
        bullets(values)
    document.add_heading("6. Estimación orientativa en horas", level=1)
    if brief.estimates:
        table(["Etapa", "Actividades principales", "Horas"], _estimate_rows(brief), _total_label(brief))
    else:
        document.add_paragraph(PENDING)
    total_min, total_max = brief.total_hours()
    document.add_paragraph(f"Total: {_hours(total_min)}–{_hours(total_max)} h. {DISCLAIMER}")
    document.add_heading("7. Próximos pasos", level=1)
    bullets(brief.next_steps)
    if brief.glossary:
        document.add_heading("Anexo A. Glosario", level=1)
        for term in brief.glossary:
            paragraph = document.add_paragraph()
            paragraph.add_run(term.term).bold = True
            paragraph.add_run(f": {term.definition}")
            if term.avoid:
                paragraph.add_run(f" Evitar: {', '.join(term.avoid)}").italic = True
    if brief.decisions:
        document.add_heading("Anexo B. Decisiones", level=1)
        for decision in brief.decisions:
            document.add_heading(f"{decision.id}. {decision.title} ({decision.status.value})", level=2)
            document.add_paragraph(decision.decision)
            if decision.considered_options:
                document.add_paragraph("Opciones consideradas: " + "; ".join(decision.considered_options))
            if decision.consequences:
                document.add_paragraph(f"Consecuencias: {decision.consequences}")
    if brief.resolved_questions:
        document.add_heading("Anexo C. Preguntas resueltas", level=1)
        table(
            ["Ronda", "Pregunta", "Respuesta"],
            [[str(q.round), f"{q.id}. {q.question}", q.answer] for q in brief.resolved_questions],
        )
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def render_all(doc: DocumentContext) -> dict[ExportFormat, bytes]:
    return {
        ExportFormat.md: render_markdown(doc).encode(),
        ExportFormat.pdf: render_pdf(doc),
        ExportFormat.docx: render_docx(doc),
    }
